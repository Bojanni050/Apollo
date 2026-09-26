"""Tests for extended document format support (.pdf, .docx, .txt) and folder browsing."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.services.documents import DOC_SUFFIXES, build_tree, list_documents, read_document
from app.services.proposals import ProposalError, plan_create, plan_edit, plan_move
from app.services.search import search_documents


def test_doc_suffixes_include_new_formats() -> None:
    assert {".txt", ".pdf", ".docx"} <= DOC_SUFFIXES


def test_read_txt_document(tmp_path: Path) -> None:
    txt_file = tmp_path / "architecture.txt"
    txt_file.write_text("Overview of the Gaia system architecture.\nDetails follow.", encoding="utf-8")

    content = read_document(tmp_path, "architecture.txt")
    assert "Overview of the Gaia system architecture." in content
    assert "Details follow." in content

    # Check tree and listing
    tree = build_tree(tmp_path)
    assert any(c.name == "architecture.txt" for c in tree.children)
    assert "architecture.txt" in list_documents(tmp_path)


def test_read_docx_document(tmp_path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Gaia Architecture Specification", level=1)
    doc.add_paragraph("This document outlines the core architecture.")
    doc.add_heading("Components", level=2)
    doc.add_paragraph("The memory component handles storage.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Memory"
    table.cell(1, 1).text = "Cache & State"

    docx_path = tmp_path / "spec.docx"
    doc.save(str(docx_path))

    content = read_document(tmp_path, "spec.docx")
    assert "# Gaia Architecture Specification" in content
    assert "## Components" in content
    assert "The memory component handles storage." in content
    assert "Memory" in content
    assert "Cache & State" in content

    # Search should also match inside the docx
    hits = search_documents(tmp_path, "storage")
    assert any(h.path == "spec.docx" for h in hits)


def test_read_pdf_document(tmp_path: Path) -> None:
    import pypdf

    writer = pypdf.PdfWriter()
    # Add a blank page with text annotation or page
    writer.add_blank_page(width=72 * 8.5, height=72 * 11)
    # Metadata title
    writer.add_metadata({"/Title": "Architecture PDF Whitepaper"})

    pdf_path = tmp_path / "whitepaper.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)

    content = read_document(tmp_path, "whitepaper.pdf")
    assert "Architecture PDF Whitepaper" in content or "Page 1" in content

    tree = build_tree(tmp_path)
    assert any(c.name == "whitepaper.pdf" for c in tree.children)


def test_proposals_binary_vs_text_rules(tmp_path: Path) -> None:
    txt = tmp_path / "guide.txt"
    txt.write_text("Step 1: Install", encoding="utf-8")

    # plan_edit on .txt works
    change = plan_edit(tmp_path, "guide.txt", "Step 1: Install & run")
    assert change.action == "edit"

    # plan_create on .txt works
    create_change = plan_create(tmp_path, "notes.txt", "Notes")
    assert create_change.action == "create"

    # plan_move on .pdf works
    pdf = tmp_path / "manual.pdf"
    pdf.write_bytes(b"%PDF-1.4 ...")
    move_change = plan_move(tmp_path, "manual.pdf", ".", new_name="renamed_manual.pdf")
    assert move_change.action == "rename"

    # plan_edit on .pdf is rejected
    with pytest.raises(ProposalError, match="Direct text editing of binary documents"):
        plan_edit(tmp_path, "manual.pdf", "new content")

    # plan_create on .docx is rejected
    with pytest.raises(ProposalError, match="Direct creation of binary documents"):
        plan_create(tmp_path, "new.docx", "new content")


def test_system_browse_folders_endpoint(client: TestClient) -> None:
    resp = client.get("/api/system/folders")
    assert resp.status_code == 200
    data = resp.json()
    assert "current_path" in data
    assert "folders" in data
    assert "drives" in data
    assert "quick_access" in data
    assert isinstance(data["folders"], list)
    assert isinstance(data["drives"], list)
    assert len(data["drives"]) > 0

    # Query with specific existing path
    curr = data["current_path"]
    resp2 = client.get(f"/api/system/folders?path={curr}")
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["current_path"] == curr
